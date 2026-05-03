from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENTATION, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
MD_PATH = OUT / "informe_servicio_maxilofacial.md"
DOCX_PATH = OUT / "informe_servicio_maxilofacial.docx"


def read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(OUT / name)


def fmt(value: object, digits: int = 1) -> str:
    if pd.isna(value):
        return "NA"
    return f"{float(value):.{digits}f}"


def fmt_p(value: object) -> str:
    if pd.isna(value):
        return "No aplicable"
    value = float(value)
    if value < 0.001:
        return "<0,001"
    return f"{value:.3f}".replace(".", ",")


def pp(value: object) -> str:
    if pd.isna(value):
        return "NA"
    return f"{100 * float(value):.1f}".replace(".", ",")


def cell_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value)


def excel_review_rows() -> list[list[str]]:
    return [
        [
            "Recidiva (Sí/No)",
            "Sustituir cada ? por Sí, No o Desconocido, usando un criterio clínico único.",
            "Permitirá estimar tasa de recidiva y comparar recidiva por abordaje con denominadores claros.",
        ],
        [
            "Tiempo Recidiva (meses)",
            "Dejar valores numéricos solo en pacientes con recidiva confirmada; en no recidiva, registrar seguimiento/censura en una columna separada.",
            "Permitirá analizar tiempo hasta recidiva o, al menos, describir seguimiento mínimo de forma interpretable.",
        ],
        [
            "Oclusión (Normal: 1/Alterada:0)",
            "Confirmar si 0 significa normal, alterada o ausencia de alteración. Ahora la columna es constante y no se puede contrastar.",
            "Evitará una conclusión errónea sobre oclusión y permitirá evaluar si hubo diferencias entre abordajes.",
        ],
        [
            "Resultado Estético (1-10)",
            "Distinguir no aplica de dato perdido; definir cuándo aplica y quién lo valoró.",
            "Permitirá comparar resultado estético solo en pacientes evaluables y con denominador honesto.",
        ],
        [
            "Complicaciones: SI: 1 /NO: 0)",
            "Resolver el registro ausente y confirmar que 1/0 significan siempre sí/no.",
            "Mejorará la comparación de seguridad y evitará excluir registros innecesariamente.",
        ],
        [
            "TIPO DE COMPLICACIÓN",
            "Separar explícitamente sin complicación de dato no registrado y normalizar categorías: parálisis, hematoma, dehiscencia, infección, otras.",
            "Permitirá resumir el perfil de complicaciones por abordaje y no solo la variable binaria.",
        ],
        [
            "Necesidad de ortognática posteriormente",
            "Resolver valores ? y estandarizar Sí/No.",
            "Permitirá valorar con más precisión la necesidad posterior de cirugía ortognática.",
        ],
        [
            "Seguimiento",
            "Añadir si es posible fecha de cirugía, fecha de última revisión y meses de seguimiento.",
            "Aportará contexto temporal a recidiva, resultado funcional y complicaciones tardías.",
        ],
    ]


def build_markdown() -> str:
    quality = read_csv("data_quality.csv")
    results = read_csv("statistical_results.csv")
    desc_cat = read_csv("descriptive_categorical.csv")

    approach_counts = (
        desc_cat[(desc_cat["group"] == "Global") & (desc_cat["variable"] == "Abordaje")]
        .sort_values("level")
        [["level", "n"]]
    )
    missing_cols = int((quality["missing_n"] > 0).sum())
    sex_values = ", ".join(
        desc_cat[(desc_cat["group"] == "Global") & (desc_cat["variable"] == "Sexo")]["level"].astype(str).tolist()
    )

    lines: list[str] = [
        "# Informe para presentación de resultados - Estudio de hiperplasia condilar",
        "",
        "Servicio de Cirugía Oral y Maxilofacial, Hospital Universitario Miguel Servet.",
        "",
        "Preparado para Juan Antonio, residente del Servicio de Cirugía Oral y Maxilofacial.",
        "",
        "## Resumen ejecutivo",
        "",
        "Se analizaron 41 registros del estudio. La unidad de análisis se considera un registro paciente/intervención. El análisis es descriptivo y comparativo exploratorio por tipo de abordaje quirúrgico: intraoral frente a preauricular.",
        "",
        "En sencillo: este documento resume los resultados principales para que Juan Antonio pueda presentarlos al servicio. No debe interpretarse como prueba causal, sino como una descripción comparativa de esta muestra.",
        "",
        "## Traducción rápida para presentar",
        "",
        "- En esta muestra, el abordaje intraoral se comportó mejor en varias variables de carga postoperatoria: menor tiempo quirúrgico, menor sangrado, menos dolor y recuperación algo más rápida.",
        "- Las diferencias más claras aparecen en sangrado, cicatriz visible y lesión del nervio facial, siempre a favor del grupo intraoral en estos datos.",
        "- El resultado funcional a 6 meses fue bueno en casi todos los casos y no mostró una diferencia relevante entre abordajes.",
        "- Las complicaciones fueron más frecuentes en el grupo preauricular, pero la incertidumbre es amplia y el resultado debe presentarse como tendencia exploratoria.",
        "- Hay que remarcar que la muestra es pequeña y observacional: los resultados orientan, pero no demuestran causalidad.",
        "",
        "En una frase: en esta serie, el abordaje intraoral parece asociarse con menor carga quirúrgica y postoperatoria, aunque la interpretación debe ser prudente por el tamaño muestral y la naturaleza retrospectiva del estudio.",
        "",
        "## Muestra y calidad de datos",
        "",
        "| Indicador | Resultado |",
        "|---|---:|",
        "| Registros analizados | 41 |",
        f"| Columnas auditadas | {len(quality)} |",
        f"| Columnas con algún dato ausente/no interpretable | {missing_cols} |",
        f"| Valores observados en Sexo | {sex_values} |",
    ]
    for _, row in approach_counts.iterrows():
        lines.append(f"| Abordaje {row['level']} | {int(row['n'])} |")

    lines.extend([
        "",
        "Advertencias de calidad relevantes:",
        "",
        "- La variable Sexo se corrigió por error de creación de la base: M/F se unificaron como Mujer y H como Hombre.",
        "- Recidiva contiene valores ?; no se estima una tasa definitiva de recidiva.",
        "- Tiempo Recidiva (meses) contiene valores no numéricos y no es analizable como tiempo hasta recodificar.",
        "- Oclusión es constante en esta base, por lo que no se contrasta inferencialmente.",
        "- Los identificadores de paciente se usaron solo para control interno de duplicados y no se incluyen en este informe.",
        "",
        "En sencillo: la base permite responder varias preguntas comparativas, pero algunas columnas necesitan aclaración antes de cerrar conclusiones definitivas.",
        "",
        "## Datos del Excel original que conviene revisar",
        "",
        "Esta es la revisión de mayor rendimiento antes de repetir o ampliar el análisis. No implica rehacer todo el Excel, sino corregir las columnas que más limitan la interpretación clínica.",
        "",
        "| Campo del Excel | Qué revisar | Qué valor aportaría al análisis |",
        "|---|---|---|",
    ])
    for field, action, value in excel_review_rows():
        lines.append(f"| {field} | {action} | {value} |")

    lines.extend([
        "",
        "En sencillo: si el servicio corrige especialmente recidiva, tiempo de recidiva, oclusión y complicaciones, el análisis ganará valor porque podrá responder preguntas clínicas que ahora solo pueden dejarse como pendientes.",
        "",
        "## Resultados comparativos principales",
        "",
        "| Desenlace | Denominador | Resultado | Efecto (IC95%) | Prueba | p | Lectura sencilla |",
        "|---|---:|---|---|---|---:|---|",
    ])

    for _, row in results[results["analysis_status"] == "Testado"].iterrows():
        if row["outcome_type"] == "cuantitativa/ordinal":
            denom = f"{int(row['n_group_1'])} vs {int(row['n_group_2'])}"
            result = f"Mediana {fmt(row['median_group_1'])} vs {fmt(row['median_group_2'])}"
            effect = f"{fmt(row['effect_value'])} ({fmt(row['ci95_low'])} a {fmt(row['ci95_high'])})"
            plain = "El abordaje intraoral presentó valores menores si el efecto es negativo."
        else:
            denom = f"{int(row['n_group_1'])} vs {int(row['n_group_2'])}"
            result = f"Eventos {int(row['events_group_1'])}/{int(row['n_group_1'])} vs {int(row['events_group_2'])}/{int(row['n_group_2'])}"
            effect = f"{pp(row['effect_value'])} pp ({pp(row['ci95_low'])} a {pp(row['ci95_high'])} pp)"
            plain = "Compara proporciones; los intervalos pueden ser amplios por eventos escasos."
        lines.append(
            f"| {row['outcome']} | {denom} | {result} | {effect} | {row['test']} | {fmt_p(row['p_value'])} | {plain} |"
        )

    skipped = results[results["analysis_status"] != "Testado"]
    lines.extend(["", "## Variables no contrastadas", ""])
    for _, row in skipped.iterrows():
        lines.append(f"- **{row['outcome']}**: {row['interpretation_note']}")

    lines.extend([
        "",
        "## Interpretación clínica prudente",
        "",
        "En esta muestra, el abordaje intraoral se asoció con menor tiempo quirúrgico, menor sangrado postoperatorio, menor dolor postoperatorio y una recuperación funcional ligeramente más corta que el abordaje preauricular. También se observaron menos lesiones de nervio facial y menos cicatrices visibles en el grupo intraoral.",
        "",
        "En sencillo: los datos favorecen al abordaje intraoral en varios resultados de carga quirúrgica y postoperatoria, pero la muestra es pequeña y observacional.",
        "",
        "La cicatriz visible debe interpretarse con especial cautela porque está muy ligada a la propia naturaleza del abordaje. No debe presentarse como efecto causal aislado sin contextualizar la técnica quirúrgica.",
        "",
        "## Métodos estadísticos",
        "",
        "- Variables cuantitativas u ordinales: mediana, IQR, media y desviación estándar como descriptiva.",
        "- Comparaciones cuantitativas entre abordajes: U de Mann-Whitney, correlación rank-biserial y diferencia de medianas con IC95% por bootstrap reproducible.",
        "- Variables binarias: prueba exacta de Fisher, diferencia de riesgos con IC95% aproximado Newcombe-Wilson y odds ratio como efecto secundario.",
        "- No se imputaron datos perdidos. Los valores ? / no aplica y textos no numéricos se documentaron y se excluyeron del análisis correspondiente.",
        "",
        "## Limitaciones",
        "",
        "- Muestra pequeña y observacional.",
        "- Múltiples comparaciones exploratorias sin endpoint primario predefinido.",
        "- Recidiva y tiempo de recidiva pendientes de aclaración.",
        "- Algunas asociaciones pueden estar condicionadas por la indicación quirúrgica y la selección de abordaje.",
        "",
        "## Archivos reproducibles asociados",
        "",
        "- `analysis_pipeline.py`",
        "- `outputs/data_quality.csv`",
        "- `outputs/codebook.csv`",
        "- `outputs/descriptive_continuous.csv`",
        "- `outputs/descriptive_categorical.csv`",
        "- `outputs/statistical_results.csv`",
    ])
    return "\n".join(lines) + "\n"


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(header_cells[i], header, bold=True, color="FFFFFF")
        shade_cell(header_cells[i], "0F766E")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    document.add_paragraph()


def add_paragraph(document: Document, text: str, style: str | None = None, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def build_docx(markdown_text: str) -> None:
    results = read_csv("statistical_results.csv")
    quality = read_csv("data_quality.csv")
    desc_cat = read_csv("descriptive_categorical.csv")

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.color.rgb = RGBColor(15, 118, 110)
    styles["Heading 2"].font.name = "Aptos"
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.color.rgb = RGBColor(51, 65, 85)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Informe para presentación de resultados\nEstudio de hiperplasia condilar")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Servicio de Cirugía Oral y Maxilofacial\nHospital Universitario Miguel Servet\nPreparado para Juan Antonio, residente del servicio")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(71, 85, 105)
    document.add_paragraph()

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Documento editable para presentación interna. No contiene identificadores de paciente ni datos fila a fila.")
    r.bold = True
    r.font.color.rgb = RGBColor(15, 118, 110)

    document.add_heading("Resumen ejecutivo", level=1)
    add_paragraph(
        document,
        "Se analizaron 41 registros del estudio. La unidad de análisis se considera un registro paciente/intervención. El análisis es descriptivo y comparativo exploratorio por tipo de abordaje quirúrgico: intraoral frente a preauricular.",
    )
    add_paragraph(
        document,
        "En sencillo: este documento resume los resultados principales para que Juan Antonio pueda presentarlos al servicio. No debe interpretarse como prueba causal, sino como una descripción comparativa de esta muestra.",
        bold_prefix="En sencillo:",
    )

    document.add_heading("Traducción rápida para presentar", level=1)
    for item in [
        "En esta muestra, el abordaje intraoral se comportó mejor en varias variables de carga postoperatoria: menor tiempo quirúrgico, menor sangrado, menos dolor y recuperación algo más rápida.",
        "Las diferencias más claras aparecen en sangrado, cicatriz visible y lesión del nervio facial, siempre a favor del grupo intraoral en estos datos.",
        "El resultado funcional a 6 meses fue bueno en casi todos los casos y no mostró una diferencia relevante entre abordajes.",
        "Las complicaciones fueron más frecuentes en el grupo preauricular, pero la incertidumbre es amplia y el resultado debe presentarse como tendencia exploratoria.",
        "Hay que remarcar que la muestra es pequeña y observacional: los resultados orientan, pero no demuestran causalidad.",
    ]:
        document.add_paragraph(item, style="List Bullet")
    add_paragraph(
        document,
        "En una frase: en esta serie, el abordaje intraoral parece asociarse con menor carga quirúrgica y postoperatoria, aunque la interpretación debe ser prudente por el tamaño muestral y la naturaleza retrospectiva del estudio.",
        bold_prefix="En una frase:",
    )

    document.add_page_break()
    document.add_heading("Muestra y calidad de datos", level=1)
    approach_counts = (
        desc_cat[(desc_cat["group"] == "Global") & (desc_cat["variable"] == "Abordaje")]
        .sort_values("level")
        [["level", "n"]]
    )
    sex_values = ", ".join(
        desc_cat[(desc_cat["group"] == "Global") & (desc_cat["variable"] == "Sexo")]["level"].astype(str).tolist()
    )
    sample_rows = [
        ["Registros analizados", "41"],
        ["Columnas auditadas", str(len(quality))],
        ["Columnas con datos ausentes/no interpretables", str(int((quality["missing_n"] > 0).sum()))],
        ["Valores observados en Sexo", sex_values],
    ]
    sample_rows.extend([[f"Abordaje {row['level']}", str(int(row["n"]))] for _, row in approach_counts.iterrows()])
    add_table(document, ["Indicador", "Resultado"], sample_rows, widths=[8.5, 6.5])

    for item in [
        "La variable Sexo se corrigió por error de creación de la base: M/F se unificaron como Mujer y H como Hombre.",
        "Recidiva contiene valores ?; no se estima una tasa definitiva de recidiva.",
        "Tiempo Recidiva (meses) contiene valores no numéricos y no es analizable como tiempo hasta recodificar.",
        "Oclusión es constante en esta base, por lo que no se contrasta inferencialmente.",
        "Los identificadores de paciente se usaron solo para control interno de duplicados y no se incluyen en este informe.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Datos del Excel original que conviene revisar", level=1)
    add_paragraph(
        document,
        "Esta es la revisión de mayor rendimiento antes de repetir o ampliar el análisis. No implica rehacer todo el Excel, sino corregir las columnas que más limitan la interpretación clínica.",
    )
    add_table(
        document,
        ["Campo del Excel", "Qué revisar", "Qué valor aportaría"],
        excel_review_rows(),
        widths=[4.8, 5.6, 5.6],
    )
    add_paragraph(
        document,
        "En sencillo: si el servicio corrige especialmente recidiva, tiempo de recidiva, oclusión y complicaciones, el análisis ganará valor porque podrá responder preguntas clínicas que ahora solo pueden dejarse como pendientes.",
        bold_prefix="En sencillo:",
    )

    result_section = document.add_section(WD_SECTION.NEW_PAGE)
    result_section.orientation = WD_ORIENTATION.LANDSCAPE
    result_section.page_width = Cm(29.7)
    result_section.page_height = Cm(21)
    result_section.top_margin = Cm(1.5)
    result_section.bottom_margin = Cm(1.5)
    result_section.left_margin = Cm(1.4)
    result_section.right_margin = Cm(1.4)
    document.add_heading("Resultados comparativos principales", level=1)
    result_rows = []
    for _, row in results[results["analysis_status"] == "Testado"].iterrows():
        if row["outcome_type"] == "cuantitativa/ordinal":
            result = f"Mediana {fmt(row['median_group_1'])} vs {fmt(row['median_group_2'])}"
            effect = f"{fmt(row['effect_value'])} (IC95% {fmt(row['ci95_low'])} a {fmt(row['ci95_high'])})"
        else:
            result = f"Eventos {int(row['events_group_1'])}/{int(row['n_group_1'])} vs {int(row['events_group_2'])}/{int(row['n_group_2'])}"
            effect = f"{pp(row['effect_value'])} pp (IC95% {pp(row['ci95_low'])} a {pp(row['ci95_high'])} pp)"
        result_rows.append([
            row["outcome"],
            f"{int(row['n_group_1'])} vs {int(row['n_group_2'])}",
            result,
            effect,
            cell_text(row["test"]),
            fmt_p(row["p_value"]),
        ])
    add_table(
        document,
        ["Desenlace", "N", "Resultado", "Efecto", "Prueba", "p"],
        result_rows,
        widths=[5.0, 1.6, 3.0, 4.0, 2.6, 1.4],
    )

    portrait_section = document.add_section(WD_SECTION.NEW_PAGE)
    portrait_section.orientation = WD_ORIENTATION.PORTRAIT
    portrait_section.page_width = Cm(21)
    portrait_section.page_height = Cm(29.7)
    portrait_section.top_margin = Cm(1.7)
    portrait_section.bottom_margin = Cm(1.7)
    portrait_section.left_margin = Cm(1.8)
    portrait_section.right_margin = Cm(1.8)
    document.add_heading("Variables no contrastadas", level=1)
    for _, row in results[results["analysis_status"] != "Testado"].iterrows():
        document.add_paragraph(f"{row['outcome']}: {row['interpretation_note']}", style="List Bullet")

    document.add_heading("Interpretación clínica prudente", level=1)
    add_paragraph(
        document,
        "En esta muestra, el abordaje intraoral se asoció con menor tiempo quirúrgico, menor sangrado postoperatorio, menor dolor postoperatorio y una recuperación funcional ligeramente más corta que el abordaje preauricular. También se observaron menos lesiones de nervio facial y menos cicatrices visibles en el grupo intraoral.",
    )
    add_paragraph(
        document,
        "En sencillo: los datos favorecen al abordaje intraoral en varios resultados de carga quirúrgica y postoperatoria, pero la muestra es pequeña y observacional.",
        bold_prefix="En sencillo:",
    )
    add_paragraph(
        document,
        "La cicatriz visible debe interpretarse con especial cautela porque está muy ligada a la propia naturaleza del abordaje. No debe presentarse como efecto causal aislado sin contextualizar la técnica quirúrgica.",
    )

    document.add_heading("Métodos estadísticos", level=1)
    for item in [
        "Variables cuantitativas u ordinales: mediana, IQR, media y desviación estándar como descriptiva.",
        "Comparaciones cuantitativas entre abordajes: U de Mann-Whitney, correlación rank-biserial y diferencia de medianas con IC95% por bootstrap reproducible.",
        "Variables binarias: prueba exacta de Fisher, diferencia de riesgos con IC95% aproximado Newcombe-Wilson y odds ratio como efecto secundario.",
        "No se imputaron datos perdidos. Los valores ? / no aplica y textos no numéricos se documentaron y se excluyeron del análisis correspondiente.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Limitaciones", level=1)
    for item in [
        "Muestra pequeña y observacional.",
        "Múltiples comparaciones exploratorias sin endpoint primario predefinido.",
        "Recidiva y tiempo de recidiva pendientes de aclaración.",
        "Algunas asociaciones pueden estar condicionadas por la indicación quirúrgica y la selección de abordaje.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "Informe generado de forma reproducible desde outputs agregados. Editable por el servicio."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(DOCX_PATH)


def main() -> None:
    markdown_text = build_markdown()
    MD_PATH.write_text(markdown_text, encoding="utf-8")
    build_docx(markdown_text)
    print(MD_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
